"""Publication figures, source-linked workbook, and academic results report."""
from pathlib import Path
import json,html
from io import BytesIO
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.colors import ListedColormap,BoundaryNorm
from matplotlib.patches import Patch
import rasterio
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment

COLORS=['#257b96','#b65a32','#577845','#775b9b']

def save(fig,out,name):
    fig.savefig(out/'figures'/f'{name}.png',dpi=240,bbox_inches='tight')
    fig.savefig(out/'figures'/f'{name}.pdf',bbox_inches='tight');plt.close(fig)

def maps(out,county,files,title,name,categorical=False):
    n=len(files);cols=2 if n==4 else min(3,n);rows=int(np.ceil(n/cols))
    fig,axes=plt.subplots(rows,cols,figsize=(4.4*cols,4.8*rows),squeeze=False)
    for ax,(label,path) in zip(axes.flat,files):
        with rasterio.open(path) as s:
            a=s.read(1,out_shape=(1100,1126),masked=True);b=s.bounds
        if categorical:
            cmap=ListedColormap(['#e2e6e6','#e69f00','#7a5195','#143e64'])
            im=ax.imshow(a,extent=(b.left,b.right,b.bottom,b.top),cmap=cmap,norm=BoundaryNorm([-.5,.5,1.5,2.5,3.5],4),interpolation='nearest')
        else:
            im=ax.imshow(a,extent=(b.left,b.right,b.bottom,b.top),cmap=ListedColormap(['#dce6e9','#197a91']),vmin=0,vmax=1,interpolation='nearest')
        county.boundary.plot(ax=ax,color='#656565',linewidth=.35)
        ax.set_title(label,fontsize=11);ax.set_aspect('equal');ax.set_xticks([]);ax.set_yticks([])
        x=b.left+12000;yy=b.bottom+13000;ax.plot([x,x+50000],[yy,yy],color='black',lw=2);ax.text(x+25000,yy+4000,'50 km',ha='center',fontsize=8)
        ax.annotate('N',xy=(.94,.92),xytext=(.94,.78),xycoords='axes fraction',ha='center',arrowprops={'arrowstyle':'->'},fontsize=9)
    for ax in list(axes.flat)[n:]:ax.set_visible(False)
    legend=[Patch(color=c,label=l) for c,l in zip(['#e2e6e6','#e69f00','#7a5195','#143e64'],['Neither','Quartile only','Loss only','Both'])] if categorical else [Patch(color='#dce6e9',label='Above annual quartile'),Patch(color='#197a91',label='At/below annual quartile')]
    fig.legend(handles=legend,loc='lower center',bbox_to_anchor=(.5,.055),ncol=len(legend),frameon=False)
    fig.suptitle(title,fontsize=16)
    fig.text(.5,.009,'30 m analysis • EPSG:5070 • White: outside valid crop support • Sources: existing yield products; UNL/FINBIN/ERS; NASS',ha='center',fontsize=8)
    fig.subplots_adjust(bottom=.17,top=.88,hspace=.17)
    save(fig,out,name)

def workbook(out):
    wb=Workbook();wb.remove(wb.active)
    names={'year_eligibility':'Year eligibility','economic_scenarios':'Economic inputs','unl_line_items':'UNL line items','nass_prices':'NASS prices','deflator':'Deflator','annual_yield':'Annual yield','sensitivity':'Sensitivity','nccpi_associations':'NCCPI uncertainty','county_statistics':'County statistics','common_valid_yield':'Common pixels','quartile_transitions':'Transitions'}
    names.update({k:v for k,v in {'unl_cost_reconciliation':'UNL reconciliation','unl_incomplete_accounts':'Incomplete UNL accounts','finbin_eligibility':'FINBIN eligibility','economic_transitions':'Economic transitions','economic_transition_gaps':'Economic gaps','economic_persistence':'Economic persistence','economic_common_valid':'Economic common pixels'}.items() if (out/f'tables/{k}.csv').exists()})
    for f,title in names.items():
        d=pd.read_csv(out/f'tables/{f}.csv');ws=wb.create_sheet(title);ws.append(d.columns.tolist())
        for row in d.itertuples(index=False,name=None):ws.append([None if pd.isna(v) else v for v in row])
        ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.fill=PatternFill('solid',fgColor='173E55');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True)
        ws.row_dimensions[1].height=44
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=min(45,max(13,len(str(col[0].value))*.8))
            for cell in list(col)[1:]:
                if isinstance(cell.value,float):cell.number_format='#,##0.00;[Red](#,##0.00);0.00'
    ws=wb.create_sheet('Budget calculator')
    ws.append(['Year','Source','Yield scenario','NASS $/bu','Cash cost $/acre','Total cost $/acre','Operator share','Mean yield Mg/ha','Revenue $/acre','Cash margin $/acre','Total return $/acre','Breakeven bu/acre','UNL eligible','Result status'])
    c=pd.read_csv(out/'tables/economic_scenarios.csv')
    baseline=pd.read_csv(out/'tables/sensitivity.csv')
    baseline=baseline[(baseline.price_factor==1)&(baseline.cost_factor==1)]
    c=c.merge(baseline[['year','source','scenario','mean_revenue_usd_ac','mean_return_usd_ac']],on=['year','source'])
    from .core import MGHA_PER_BUAC
    for i,r in enumerate(c.to_dict('records'),2):
        mean_yield=r['mean_revenue_usd_ac']/(r['nass_price_usd_bu']*r['operator_share'])*MGHA_PER_BUAC
        assert abs(mean_yield/MGHA_PER_BUAC*r['nass_price_usd_bu']*r['operator_share']-r['total_cost_usd_ac']-r['mean_return_usd_ac'])<1e-7
        ws.append([r['year'],r['source'],r['scenario'],r['nass_price_usd_bu'],None if pd.isna(r['cash_cost_usd_ac']) else r['cash_cost_usd_ac'],r['total_cost_usd_ac'],r['operator_share'],mean_yield,
                   f'=H{i}/{MGHA_PER_BUAC}*D{i}*G{i}',f'=IF(ISNUMBER(E{i}),I{i}-E{i},"")',f'=I{i}-F{i}',f'=F{i}/(D{i}*G{i})',True,'EXPERIMENTAL 2021' if r['year']==2021 else 'Primary '+str(r['year'])])
        for col in range(4,13):ws.cell(i,col).number_format='#,##0.00;[Red](#,##0.00);0.00'
    for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='173E55');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True)
    ws.row_dimensions[1].height=40
    ws.freeze_panes='C2'
    for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=24
    ws=wb.create_sheet('Read me');
    for text in ['CH4 exact-year UNL gated research budgets','All currency USD. Blank cells are unavailable, never zero.',
        'Original-year rotation-priority approximations are labeled in Year eligibility. 2021 is experimental and excluded from temporal summaries.',
        'FINBIN returns are operator-account proxies, not complete economic profit. ERS uses planted acres and all practices.',
        '2019 UNL printed cash cost 2.14/bu does not reconcile; derived cash cost is 363.18/acre.',
        'Budget calculator uses each source-domain valid-area mean yield; mapped loss areas use individual unchanged pixels.',
        'Formula values recalculate when opened in Excel; verified numeric results are in the other sheets.',
        'Source originals, checksums, pages, definitions and exceptions are provided in sources and the report.']:ws.append([text])
    ws.column_dimensions['A'].width=125
    wb.move_sheet(wb['Budget calculator'],offset=-len(wb.worksheets)+2)
    buffer=BytesIO();wb.save(buffer);buffer.seek(0)
    check=load_workbook(buffer)
    assert check['Year eligibility'].max_row==12 and check['Budget calculator']['I2'].data_type=='f' and check['Budget calculator'].max_row==len(c)+1
    (out/'CH4_budget_workbook.xlsx').write_bytes(buffer.getvalue())

def create_report(out,county,root):
    out=Path(out);a=pd.read_csv(out/'tables/annual_yield.csv');s=pd.read_csv(out/'tables/sensitivity.csv');c=pd.read_csv(out/'tables/county_statistics.csv');n=pd.read_csv(out/'tables/nccpi_associations.csv')
    factors=pd.read_csv(out/'tables/economic_scenarios.csv').set_index(['year','source']).to_2021_dollars
    adjustment=np.array([factors.loc[(r.year,r.source)] for r in s.itertuples()])
    for col in ['mean_revenue_usd_ac','mean_cash_margin_usd_ac','return_p05','return_median','return_p95']:
        s[col+'_2021dollars']=s[col]*adjustment
    s.to_csv(out/'tables/sensitivity.csv',index=False)
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False})
    fig,ax=plt.subplots(2,1,figsize=(10,7),sharex=True)
    for color,(scenario,d) in zip(COLORS,a.groupby('scenario')):
        primary=d[d.year<2021];ax[0].plot(primary.year,primary['mean'],'o-',color=color,label=scenario)
        ax[1].plot(primary.year,primary.quartile_ha/1000,'o-',color=color,label=scenario)
    ax[0].set_ylabel('Mean yield (Mg/ha)');ax[1].set_ylabel('Quartile-marginal area (1,000 ha)');ax[1].set_xticks(range(2001,2020,2));ax[0].legend();ax[0].set_title('Primary annual footprints, 2001–2019; 2019 mask source changed')
    save(fig,out,'annual_yield_and_quartile')
    fig,axes=plt.subplots(2,1,figsize=(11,8),sharex=True)
    for ax,(scenario,d),color in zip(axes,a[a.year<2021].groupby('scenario'),COLORS):
        boxes=[dict(label=str(r.year),med=r.median,q1=r.p25,q3=r.p75,whislo=r.p05,whishi=r.p95,fliers=[]) for r in d.itertuples()]
        ax.bxp(boxes,patch_artist=True,boxprops={'facecolor':color,'alpha':.6},showfliers=False)
        ax.set_ylabel('Yield (Mg/ha)');ax.set_title(scenario+' • pixel-area distribution')
    axes[-1].set_xlabel('Odd production year; 2019 mask source changed')
    fig.suptitle('Primary yield distributions: box = 25th–75th percentiles; whiskers = 5th–95th')
    fig.tight_layout();save(fig,out,'primary_yield_distributions')
    base=s[(s.price_factor==1)&(s.cost_factor==1)]
    for year in sorted(base.year.unique()):
        d=base[(base.year==year)&(base.scenario=='M1_fixed')]
        fig,ax=plt.subplots(figsize=(9,5));x=np.arange(len(d));ax.bar(x-.18,d.loss_ha/1000,width=.36,label='Negative source-account return',color='#775b9b');ax.bar(x+.18,d.quartile_ha/1000,width=.36,label='Lowest yield quartile',color='#197a91')
        ax.set_xticks(x,d.source);ax.set_ylabel('Area (1,000 ha)');ax.legend();ax.set_title(f'{year} M1 baseline'+(' — EXPERIMENTAL' if year==2021 else ' — 2019 footprint caveat'))
        fig.text(.02,.01,'FINBIN county restricted to four county–AOI intersections; FINBIN excludes unestablished opportunity costs.',fontsize=8)
        save(fig,out,f'economic_quartile_{year}')
        d=s[(s.year==year)&(s.scenario=='M1_fixed')&(s.source=='UNL')]
        matrix=d.pivot(index='cost_factor',columns='price_factor',values='loss_ha')/1000
        fig,ax=plt.subplots(figsize=(6,5));im=ax.imshow(matrix,origin='lower',cmap='Blues')
        for i in range(3):
            for j in range(3):ax.text(j,i,f'{matrix.iloc[i,j]:.1f}',ha='center',va='center')
        ax.set_xticks(range(3),['−15%','Baseline','+15%']);ax.set_yticks(range(3),['−15%','Baseline','+15%']);ax.set_xlabel('Price');ax.set_ylabel('Cost');ax.set_title(f'{year} UNL M1 loss area (1,000 ha)'+('\nExperimental' if year==2021 else ''));save(fig,out,f'sensitivity_{year}')
    fig,ax=plt.subplots(figsize=(11,6))
    for name,d in c[(c.source=='yield_only')&(c.scenario=='M1_fixed')&(c.year<2021)].groupby('NAME'):
        ax.plot(d.year,100*d.quartile_ha/d.valid_ha,'o-',ms=3,label=name,alpha=.8)
    ax.set_ylabel('Annual quartile-marginal share (%)');ax.set_xticks(range(2001,2020,2));ax.set_title('County–AOI trajectories on changing valid corn footprints');ax.legend(bbox_to_anchor=(1.01,1),fontsize=8);save(fig,out,'county_quartile_trajectories')
    fig,ax=plt.subplots(figsize=(10,5));d=n[(n.source=='yield_only')&(n.metric=='quartile_marginality')&(n.block_m==10000)&(n.year<2021)]
    for color,(scenario,g) in zip(COLORS,d.groupby('scenario')):
        ax.plot(g.year,g.r,'o-',color=color,label=scenario)
        ax.vlines(g.year,g.ci_low,g.ci_high,color=color)
        ax.hlines(g.ci_low,g.year-.1,g.year+.1,color=color)
        ax.hlines(g.ci_high,g.year-.1,g.year+.1,color=color)
    ax.axhline(0,color='gray',lw=.7);ax.set_ylabel('NCCPI vs quartile indicator: Pearson r');ax.set_title('Spatial block uncertainty: 10 km blocks, 1,999 resamples');ax.legend();save(fig,out,'nccpi_association')
    maps(out,county,[(str(y),out/f'rasters/quartile_{y}_M1_fixed.tif') for y in range(2001,2020,2)],'Lowest yield quartile • M1 • primary years','primary_quartile_maps')
    for year in sorted(base.year.unique()):
        maps(out,county,[(source.replace('_',' '),out/f'rasters/overlap_{year}_M1_fixed_{source}.tif') for source in base[base.year==year].source.unique()],f'{year} M1 marginality overlap'+(' • EXPERIMENTAL' if year==2021 else ' • mask source changed'),'overlap_maps_'+str(year),True)
    workbook(out)
    eligibility=pd.read_csv(out/'tables/year_eligibility.csv')
    methods=(HERE/'METHODS.md').read_text(encoding='utf8')
    u=base[(base.year==2019)&(base.source=='UNL')]
    u21=base[(base.year==2021)&(base.source=='UNL')&(base.scenario=='M1_fixed')].iloc[0]
    finding=' '.join(f"Under the 2019 UNL baseline, {r.scenario} classifies {r.loss_ha:,.2f} ha ({100*r.loss_ha/r.valid_ha:.2f}% of valid crop area) as economic loss, compared with {r.quartile_ha:,.2f} ha in the lowest yield quartile. The definitions overlap on {r.both_ha:,.2f} ha and disagree on {r.disagree_ha:,.2f} ha; mean modeled return is {'−' if r.mean_return_usd_ac<0 else ''}${abs(r.mean_return_usd_ac):,.2f}/acre in nominal dollars." for r in u.itertuples())
    common=pd.read_csv(out/'tables/common_valid_yield.csv')
    common_text=' '.join(f"{scenario} has {g.common_ha.iloc[0]:,.2f} ha valid in every primary observation; its common-domain annual mean yields range from {g['mean'].min():.2f} to {g['mean'].max():.2f} Mg/ha." for scenario,g in common.groupby('scenario'))
    soil=n[(n.source=='yield_only')&(n.metric=='quartile_marginality')&(n.block_m==10000)&(n.year<2021)]
    soil_text=f"Primary annual NCCPI–quartile indicator correlations range from {soil.r.min():.3f} to {soil.r.max():.3f}. The accompanying table and chart report 10 km block intervals, with 5 km and 20 km sensitivity in the source table. These describe association with modeled yield ranks, not observed producer profitability."
    lines=['# CH4 marginality results', '', 'Analysis of modeled dryland corn in Nebraska MLRA 106. This is a modeled enterprise-budget assessment, not independent profitability validation.',
        '',f"Primary economic calculations are eligible only for 2019. Experimental 2021 is separate. Blocked years: {', '.join(map(str,eligibility[~eligibility.eligible].year))}.",
        '', '## Principal results',finding,'',common_text,'',soil_text,
        '', '## Experimental 2021',f"The experimental UNL M1 baseline produces {u21.loss_ha:,.2f} ha of modeled economic loss and {u21.quartile_ha:,.2f} ha of quartile marginality over {u21.valid_ha:,.2f} valid hectares. M2 is identical because measured 2021 HI is unavailable. Differences from 2019 combine changed prices, budgets, footprints and experimental productivity inputs; they are not interpreted as a temporal economic trend.",
        '', '## Annual yield and quartile marginality',a[['year','scenario','experimental','mean','valid_ha','missing_ha','cutoff_Mg_ha','quartile_ha','quartile_percent']].to_markdown(index=False,floatfmt='.3f'),
        '', '## Economic and accounting-proxy results',base[['year','scenario','source','full_economic_account','valid_ha','loss_ha','both_ha','disagree_ha','mean_return_usd_ac','mean_return_2021usd_ac']].to_markdown(index=False,floatfmt='.2f'),
        '', 'Only one primary year passes the UNL gate, so primary economic persistence and economic transitions cannot be estimated. Quartile persistence and transitions use 2001–2019 valid corn observations. Experimental 2021 is excluded.',
        '', '## Methods and limitations',methods]
    text='\n'.join(lines);(out/'CH4_methods_results.md').write_text(text,encoding='utf8')
    import markdown
    body=markdown.markdown(text,extensions=['tables'])
    for p in sorted((out/'figures').glob('*.png')):body+=f'<figure><img src="figures/{p.name}" style="max-width:100%"><figcaption>{p.stem.replace("_"," ")}</figcaption></figure>'
    (out/'CH4_methods_results.html').write_text('<!doctype html><meta charset="utf-8"><title>CH4 marginality</title><style>body{max-width:1200px;margin:50px auto;font:17px/1.6 Georgia;color:#163447}table{font:12px/1.4 Arial;border-collapse:collapse;width:100%}td,th{padding:5px;border-bottom:1px solid #ddd}h1,h2,h3{font-family:Arial}a{color:#197a91}</style>'+body,encoding='utf8')

from .prepare_budgets import HERE
